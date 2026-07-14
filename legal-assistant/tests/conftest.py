import json
import os
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

# 代码平铺于 scripts/，按 pyproject 的 package-dir 映射为 legal_assistant 包；
# 未 pip install 时由此 shim 注册包，保证测试可独立运行
if "legal_assistant" not in sys.modules:
    import importlib.util

    _scripts = Path(__file__).resolve().parent.parent / "scripts"
    _spec = importlib.util.spec_from_file_location(
        "legal_assistant", _scripts / "__init__.py",
        submodule_search_locations=[str(_scripts)],
    )
    _pkg = importlib.util.module_from_spec(_spec)
    sys.modules["legal_assistant"] = _pkg
    _spec.loader.exec_module(_pkg)

from legal_assistant.config import load_config  # noqa: E402


class FakeLark:
    """内存版 LarkClient：接口一致，记录写入，支持 == 过滤。"""

    def __init__(self):
        self.tables = {"reimburse": [], "contracts": [], "output_invoices": []}
        self.messages = []
        self.attachments = []
        self._next_id = 1
        self.dry_run = False
        self.skipped_writes = []

    # ---- 读 ----
    def record_list(self, table_key, field_names=None, filter_json=None):
        records = self.tables[table_key]
        if filter_json:
            for field, op, value in filter_json.get("conditions", []):
                assert op == "=="
                records = [r for r in records if str(r["fields"].get(field, "")) == str(value)]
        return [dict(r) for r in records]

    def find_by_field(self, table_key, field_name, value):
        return self.record_list(
            table_key, filter_json={"logic": "and", "conditions": [[field_name, "==", value]]}
        )

    # ---- 写 ----
    def record_create(self, table_key, fields):
        if self.dry_run:
            self.skipped_writes.append((table_key, fields))
            return None
        record_id = f"rec_{self._next_id}"
        self._next_id += 1
        self.tables[table_key].append({"record_id": record_id, "fields": dict(fields)})
        return record_id

    def record_update(self, table_key, record_id, fields):
        if self.dry_run:
            self.skipped_writes.append((table_key, record_id, fields))
            return
        for rec in self.tables[table_key]:
            if rec["record_id"] == record_id:
                rec["fields"].update(fields)
                return
        raise ValueError(f"record {record_id} not found")

    def upload_attachment(self, table_key, record_id, field_name, file_path):
        self.attachments.append((table_key, record_id, field_name, str(file_path)))

    def send_markdown(self, user_id, markdown, idempotency_key=None):
        self.messages.append({"user_id": user_id, "markdown": markdown, "key": idempotency_key})


CONFIG_TEMPLATE = """
firm_name: "北京市隆安（深圳）律师事务所"
paths:
  invoice_inbox: "{root}/发票收件箱"
  invoice_output: "{root}/发票整理"
  invoice_archive: "{root}/发票收件箱/已处理"
  output_invoice_inbox: "{root}/开票收件箱"
  output_invoice_archive: "{root}/开票收件箱/已处理"
  scan_inbox: "{root}/扫描件"
  scan_output: "{root}/扫描件"
  state_dir: "{root}/state"
  journal_dir: "{root}/journal"
lark:
  bin: "lark-cli"
  base_token: "bascnTEST"
  identity: "user"
  tables:
    reimburse: "tblA"
    contracts: "tblB"
    output_invoices: "tblC"
  push:
    user_id: "ou_test"
fields:
  reimburse:
    invoice_no: "发票号码"
    invoice_type: "发票类型"
    amount: "金额"
    payee: "收款方"
    invoice_date: "开票日期"
  output_invoices:
    invoice_no: "发票号"
    member: "团队成员"
    contract_no: "合同号"
    payer: "付款方"
    invoice_date: "开票日期"
    amount: "发票金额"
    contract_link: "合同台账"
  contracts:
    contract_no: "律所合同号"
    client: "客户名称"
    invoiced: "已开票"
    credit_code: "信用代码/身份证"
    amount: "合同金额(元)"
    sign_date: "签约日期"
    attachment: "合同附件"
invoice:
  zip_stable_seconds: 1
scan:
  categories:
    合同文件: ["合同", "协议"]
    诉讼函件: ["律师函", "判决"]
    财务相关: ["发票", "dzfp"]
  trash_keywords: ["test", "副本"]
  review_dir_name: "待人工复核"
  trash_dir_name: "临时垃圾箱"
  ocr_command: ""
weekly:
  xlsx_output_dir: "{root}/周报"
"""


@pytest.fixture
def cfg(tmp_path):
    config_path = tmp_path / "config.yaml"
    config_path.write_text(CONFIG_TEMPLATE.format(root=tmp_path), encoding="utf-8")
    c = load_config(config_path)
    for key in ("invoice_inbox", "output_invoice_inbox", "scan_inbox"):
        c.paths[key].mkdir(parents=True, exist_ok=True)
    return c


@pytest.fixture
def fake_lark():
    return FakeLark()


# QQ 导出 xlsx 的 12 列表头（与真实样例一致）
QQ_HEADER = [
    "备注", "发票类型", "开票日期", "付款方", "付款方纳税人识别号", "发票代码",
    "发票号码", "发票校验码", "收款方", "税前金额", "税额", "金额",
]


def make_raw_zip(path: Path, entries):
    """手工构造 ZIP（store、无 UTF-8 flag），复现 QQ/Windows 导出的 GBK 文件名条目。
    Python zipfile 对非 ASCII 名强制加 UTF-8 flag，只能裸写二进制。
    entries: [(name_bytes, data_bytes)]"""
    import binascii
    import struct

    blob = b""
    central = b""
    offsets = []
    for name, data in entries:
        offsets.append(len(blob))
        crc = binascii.crc32(data) & 0xFFFFFFFF
        # local file header：version=20, flags=0（无 0x800）, method=0(store)
        blob += struct.pack("<IHHHHHIIIHH", 0x04034B50, 20, 0, 0, 0, 0x21,
                            crc, len(data), len(data), len(name), 0)
        blob += name + data
    for (name, data), offset in zip(entries, offsets):
        crc = binascii.crc32(data) & 0xFFFFFFFF
        central += struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, 20, 20, 0, 0, 0, 0x21,
                               crc, len(data), len(data), len(name), 0, 0, 0, 0, 0, offset)
        central += name
    eocd = struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, len(entries), len(entries),
                       len(central), len(blob), 0)
    path.write_bytes(blob + central + eocd)
    old = time.time() - 3600
    os.utime(str(path), (old, old))
    return path


def make_invoice_zip(path: Path, rows, pdf_names=None, include_xlsx=True):
    """构造 QQ 邮箱发票导出 zip。rows: [(类型, 日期, 发票号, 收款方, 金额)]。"""
    with zipfile.ZipFile(str(path), "w") as zf:
        if include_xlsx:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(QQ_HEADER)
            total = 0.0
            for inv_type, inv_date, inv_no, payee, amount in rows:
                ws.append(["", inv_type, inv_date, "北京市隆安（深圳）律师事务所", "TAX123",
                           "", inv_no, "", payee, amount, 0, amount])
                total += amount
            ws.append(["总计", "", "", "", "", "", "", "", "", "", "", total])
            xlsx_tmp = path.parent / "_stats_tmp.xlsx"
            wb.save(str(xlsx_tmp))
            zf.write(str(xlsx_tmp), "发票统计.xlsx")
            xlsx_tmp.unlink()
        for name in pdf_names or []:
            zf.writestr(name, b"%PDF-1.4 fake")
    # mtime 拉老，通过 stable_seconds 检查
    old = time.time() - 3600
    os.utime(str(path), (old, old))
    return path
