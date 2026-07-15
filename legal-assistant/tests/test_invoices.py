"""Slice 1：进项发票流程。"""

from legal_assistant.invoices import run_invoice_once
from legal_assistant.journal import NullJournal
from legal_assistant.state import Ledger, FAILED

from conftest import make_invoice_zip

ROWS = [
    ("餐饮", "2026/07/04", "26100000000000000001", "示例餐饮有限公司", 378.00),
    ("服务", "2026/07/02", "26100000000000000002", "示例咨询服务有限公司", 300.00),
]
PDFS = [
    "260704_378.00_示例餐饮有限公司.pdf",
    "260702_300.00_示例咨询服务有限公司.pdf",
]


def test_full_flow(cfg, fake_lark):
    zip_path = make_invoice_zip(cfg.paths["invoice_inbox"] / "QQ导出.zip", ROWS, PDFS)
    results = run_invoice_once(cfg, fake_lark, NullJournal())
    assert len(results) == 1 and results[0].ok
    assert len(results[0].created) == 2
    # 飞书写入 5 字段
    fields = fake_lark.tables["reimburse"][0]["fields"]
    assert fields["发票号码"] == "26100000000000000001"
    assert fields["发票类型"] == "餐饮"
    assert fields["金额"] == 378.00
    # PDF 按类型归档
    out_dirs = {p.name for p in cfg.paths["invoice_output"].rglob("*") if p.is_dir()}
    assert {"餐饮", "服务"} <= out_dirs
    # zip 移入已处理，收件箱清空
    assert not zip_path.exists()
    assert (cfg.paths["invoice_archive"] / "QQ导出.zip").exists()


def test_rerun_is_idempotent(cfg, fake_lark):
    make_invoice_zip(cfg.paths["invoice_inbox"] / "a.zip", ROWS, PDFS)
    run_invoice_once(cfg, fake_lark, NullJournal())
    assert len(fake_lark.tables["reimburse"]) == 2
    # 同样内容再来一个 zip（模拟重复导出）：远端去重，不重复写
    make_invoice_zip(cfg.paths["invoice_inbox"] / "b.zip", ROWS, PDFS)
    results = run_invoice_once(cfg, fake_lark, NullJournal())
    assert len(fake_lark.tables["reimburse"]) == 2
    assert len(results[0].duplicates) == 2
    assert len(results[0].created) == 0


def test_missing_pdf_still_uploads(cfg, fake_lark):
    make_invoice_zip(cfg.paths["invoice_inbox"] / "a.zip", ROWS, PDFS[:1])
    results = run_invoice_once(cfg, fake_lark, NullJournal())
    r = results[0]
    assert r.ok and len(r.created) == 2
    assert r.missing_pdf == ["26100000000000000002"]


def test_orphan_pdf_goes_to_review(cfg, fake_lark):
    pdfs = PDFS + ["260101_999.00_不认识的公司.pdf"]
    make_invoice_zip(cfg.paths["invoice_inbox"] / "a.zip", ROWS, pdfs)
    results = run_invoice_once(cfg, fake_lark, NullJournal())
    assert results[0].orphan_pdfs == ["260101_999.00_不认识的公司.pdf"]
    review = list(cfg.paths["invoice_output"].rglob("待人工复核/*.pdf"))
    assert len(review) == 1


def test_bad_zip_marked_failed_and_stays(cfg, fake_lark):
    # 只有 PDF 没有 xlsx 的半成品
    zip_path = make_invoice_zip(
        cfg.paths["invoice_inbox"] / "half.zip", [], PDFS, include_xlsx=False
    )
    results = run_invoice_once(cfg, fake_lark, NullJournal())
    assert not results[0].ok
    assert zip_path.exists()  # 留原地
    ledger = Ledger(cfg.paths["state_dir"], "invoices")
    assert ledger.status("half.zip") == FAILED
    assert not fake_lark.tables["reimburse"]


def test_dry_run_touches_nothing(cfg, fake_lark):
    fake_lark.dry_run = True
    zip_path = make_invoice_zip(cfg.paths["invoice_inbox"] / "a.zip", ROWS, PDFS)
    run_invoice_once(cfg, fake_lark, NullJournal(), dry_run=True)
    assert zip_path.exists()  # zip 未移动
    assert not fake_lark.tables["reimburse"]  # 未写入
    assert len(fake_lark.skipped_writes) == 2  # 但记录了将要写什么
    assert not (cfg.paths["invoice_output"]).exists() or not list(
        cfg.paths["invoice_output"].rglob("*.pdf")
    )


def test_repair_zip_name_encodings():
    import zipfile

    from legal_assistant.invoices import repair_zip_name

    orig = "260704_378.00_示例餐饮有限公司.pdf"
    for enc in ("utf-8", "gbk"):
        info = zipfile.ZipInfo(orig.encode(enc).decode("cp437"))
        info.flag_bits = 0  # 模拟无 UTF-8 flag 的 QQ/Windows 导出
        assert repair_zip_name(info) == orig
    # 带 flag 的正常条目原样返回
    info = zipfile.ZipInfo(orig)
    info.flag_bits = 0x800
    assert repair_zip_name(info) == orig


def test_gbk_zip_full_flow(cfg, fake_lark):
    """真实 QQ/Windows 导出复现：无 UTF-8 flag、GBK 文件名的 zip 全流程。"""
    import io

    import openpyxl

    from conftest import QQ_HEADER, make_raw_zip

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(QQ_HEADER)
    ws.append(["", "餐饮", "2026/07/04", "示例律师事务所", "TAX", "",
               "26100000000000000009", "", "示例餐饮有限公司", 356.60, 21.40, 378.00])
    buf = io.BytesIO()
    wb.save(buf)

    pdf_name = "260704_378.00_示例餐饮有限公司.pdf".encode("gbk")
    make_raw_zip(
        cfg.paths["invoice_inbox"] / "qq_gbk.zip",
        [("发票统计.xlsx".encode("gbk"), buf.getvalue()),
         (pdf_name, b"%PDF-1.4 fake")],
    )
    results = run_invoice_once(cfg, fake_lark, NullJournal())
    r = results[0]
    assert r.ok and len(r.created) == 1
    assert not r.missing_pdf and not r.orphan_pdfs  # 文件名修复后 join 成功
    # 归档文件名已还原为正确中文
    archived = list(cfg.paths["invoice_output"].rglob("餐饮/*.pdf"))
    assert len(archived) == 1
    assert archived[0].name == "260704_378.00_示例餐饮有限公司.pdf"


def test_fresh_zip_waits_for_stability(cfg, fake_lark, tmp_path):
    import zipfile

    # 刚落盘的 zip（mtime = now）不处理
    fresh = cfg.paths["invoice_inbox"] / "fresh.zip"
    with zipfile.ZipFile(str(fresh), "w") as zf:
        zf.writestr("x.txt", "downloading")
    cfg2 = cfg
    cfg2.zip_stable_seconds = 3600
    results = run_invoice_once(cfg2, fake_lark, NullJournal())
    assert results == []
